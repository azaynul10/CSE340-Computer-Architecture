"""
Builds the CSE340 Problem 15 report as a Word .docx with the exact required
formatting: Times New Roman 11, single column, single spacing.

Run:  python build_docx.py
Out:  weather-station-report.docx  (rename to <ID>_<SECTION>_<Name>.docx before submitting)
"""

import os

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

OUT = "weather-station-report.docx"

TITLE = "Scaling the Architecture of a Village Weather Station Over Time"

INTRO = (
    "A weather station built for a rural school is a useful lens on architectural scaling because its "
    "constraints are absolute rather than negotiable: a fixed energy envelope, a fixed budget, and no "
    "technician on site. The original deployment logs two sensors hourly on an 8-bit microcontroller with "
    "2 KB of SRAM. Two years later the community wants ten times the sampling rate, real-time storm "
    "alerting, monthly Bluetooth synchronisation and an on-device chart display, on the same power budget. "
    "This report argues that these requests cannot be met by adding features to the existing design, because "
    "the limiting resources are SRAM capacity and the single-threaded control flow rather than clock speed. "
    "It then proposes a revised memory hierarchy and concurrency model, and quantifies the result against a "
    "budget of 100 mAh per day."
)

BACKGROUND = [
    (
        "The original platform is a classic 8-bit AVR-class device: 16 MHz, 2 KB SRAM, 32 KB flash, no cache "
        "and no operating system (Microchip Technology, 2016). Memory is flat and uniformly accessed, so "
        "every load takes a fixed number of cycles. Patterson and Hennessy (2021, Chapter 5) frame the "
        "memory hierarchy as "
        "a trade of capacity against latency, exploiting locality so that a large memory approximates the "
        "speed of a small one. This part sits at a degenerate point in that hierarchy: with one level there "
        "is no locality to exploit and no cache-miss variability. The result is complete timing determinism, "
        "which is why such parts persist in control applications."
    ),
    (
        "Three documented facts about the requested features determine what breaks. First, the SD card "
        "interface is block-oriented: the physical layer specification fixes a 512-byte block as the unit of "
        "transfer (SD Association, 2010), so any driver must hold a 512-byte buffer, which is a quarter of "
        "the available SRAM before application code exists. Second, a 1.8-inch display of the common ST7735 "
        "class drives a 128 by 160 pixel panel at 16 bits per pixel (Sitronix Technology, 2011). A full "
        "frame buffer is "
        "therefore 128 x 160 x 2 = 40 KB, twenty times the entire SRAM of the original part; even a "
        "one-bit-per-pixel monochrome buffer needs 2.5 KB and still does not fit. Third, Bluetooth Classic "
        "modules of the HC-05 class draw tens of milliamperes continuously while connectable, whereas "
        "Bluetooth Low Energy is designed around short connection events with microampere idle current "
        "(Bluetooth Special Interest Group, 2019)."
    ),
    (
        "Prior work on energy-constrained sensing converges on two strategies. Duty cycling with aggressive "
        "sleep states is standard, since sleep currents sit three to four orders of magnitude below active "
        "current (Espressif Systems, 2023; Microchip Technology, 2016), and modern sensors support it by "
        "drawing microamperes between commanded measurements (Sensirion AG, 2019). Batching writes is the "
        "second: NOR flash is programmed in pages and erased in sectors with a finite endurance of roughly "
        "100,000 cycles per sector (Winbond Electronics, 2021), so aggregating small records into one "
        "page-aligned write reduces both energy per byte and wear. This is the spatial-locality argument "
        "that motivates cache line fills in Patterson and Hennessy (2021, Chapter 5), applied to a "
        "non-volatile level of the hierarchy. The concurrency literature treats a blocking write inside a "
        "single-threaded loop as a latency hazard; decoupling producers from consumers through a ring "
        "buffer, or scheduling them as priority tasks under a small real-time kernel, is the conventional "
        "remedy (Barry, 2016; Patterson & Hennessy, 2021, Chapter 6)."
    ),
]

METH_1 = (
    "Step 1 - establish where the original design fails. Table 1 costs the naive approach of attaching the "
    "new peripherals to the existing superloop."
)

TABLE1_CAP = "Table 1: Daily energy cost of adding the requested features without architectural change"
TABLE1 = [
    ("Component", "Assumption", "mAh/day"),
    ("MCU active continuously, never sleeps", "12 mA x 24 h", "288"),
    ("Bluetooth Classic module always connectable", "30 mA x 24 h", "720"),
    ("Display and backlight always on", "30 mA x 24 h", "720"),
    ("SD card idle current plus 240 writes", "~0.5 mA average", "12"),
    ("Total", "", "~1740"),
]

METH_2 = (
    "Against a 100 mAh per day budget this is seventeen times over, and the memory position is not a matter "
    "of degree: the 512-byte SD buffer, a serial receive buffer and any display buffer cannot coexist in "
    "2 KB. Sampling every six minutes also produces 240 records per day, needing 1200 bytes of buffering at "
    "five bytes each. These are hard limits, so the platform itself must change."
)

METH_3 = (
    "Step 2 - choose a processor class and memory hierarchy. I select a 32-bit RISC-V microcontroller with an "
    "integrated 2.4 GHz radio, of the ESP32-C3 class: a 160 MHz core, 384 KB of ROM, 400 KB of on-chip SRAM "
    "of which 16 KB is allocated to cache, native Bluetooth Low Energy 5.0, and a deep-sleep current of 5 uA "
    "(Espressif Systems, 2023). This resolves all three memory limits at once and removes the external "
    "Bluetooth module entirely. Selecting RISC-V also keeps the instruction-set model consistent with the "
    "course reference."
)

METH_4 = (
    "The cache deserves explicit treatment, because it reintroduces a trade-off the 8-bit part did not have. "
    "Espressif documents flash access on this part as accelerated by cache, which is necessary to hide flash "
    "latency at 160 MHz, but a cache miss makes execution time variable, and variable latency is precisely "
    "what a storm-alert path must not have. The design therefore partitions code by requirement: the alert "
    "handler and sampling routine are placed in on-chip SRAM so their timing is deterministic, while display "
    "rendering and synchronisation, which are throughput-oriented and not deadline-bound, execute from "
    "cached flash. This is the central design decision in this proposal."
)

METH_5 = (
    "Step 3 - restructure the control flow. The superloop is replaced by three cooperating elements. A "
    "hardware timer interrupt samples the sensors and writes fixed-size records into an SRAM ring buffer, so "
    "acquisition cannot be delayed by slower activity. A low-priority task drains that buffer to NOR flash "
    "every four hours in page-aligned batches. Storm detection becomes event-driven rather than polled: the "
    "threshold comparison runs inside the timer handler and asserts the buzzer directly, so alert latency is "
    "bounded by the sampling interval rather than by the main loop. The display and monthly synchronisation "
    "are on-demand consumers, triggered by a button press and a scheduled wake."
)

METH_6 = "Step 4 - verify the budget on the same assumptions."

TABLE2_CAP = "Table 2: Daily energy cost of the revised event-driven design"
TABLE2 = [
    ("Component", "Assumption", "mAh/day"),
    ("Deep sleep for approximately 23.9 h", "5 uA", "0.12"),
    ("240 sensor wakes at 50 ms each", "20 mA active", "0.07"),
    ("Flash batch writes, six per day", "15 mA for ~3 s total", "0.01"),
    ("Storm alerting, two minutes average", "25 mA", "0.83"),
    ("Display on demand, 10 views x 20 s", "40 mA", "2.20"),
    ("Monthly BLE sync, 60 s at 100 mA, amortised", "", "0.06"),
    ("Total", "", "~3.3"),
]

DISCUSSION = [
    (
        "The revised design consumes about 3.3 per cent of the daily budget. The dominant remaining cost is "
        "the display backlight rather than computation, which is the main design lesson: once sleep states "
        "and batching are applied, the processor stops being the constraint. A memory-in-pixel or "
        "electrophoretic panel would cut that 2.20 mAh to roughly a tenth because it retains an image "
        "without power, but at 900 to 1500 taka it strains the original sub-1200 taka budget, so the honest "
        "recommendation is an on-demand backlight now and an e-paper upgrade only if the budget grows. Two "
        "limitations should be stated plainly: every figure is derived from datasheet typical values rather "
        "than measured on hardware, and solar harvesting varies sharply with monsoon cloud cover, so margin "
        "of this order is a requirement rather than a luxury."
    ),
    (
        "The sustainability argument follows from the same numbers. Provisioning 400 KB of SRAM where the "
        "application needs perhaps 20 KB costs little at purchase but avoids replacing an entire board when "
        "requirements grow, which is precisely the failure this scenario describes; discarded electronic "
        "equipment is already the fastest-growing domestic waste stream, and design decisions that extend "
        "service life are among the few available levers on it (Forti et al., 2020). Page-aligned batching "
        "reduces flash wear by writing whole pages instead of five-byte records. Keeping sensing, compute and "
        "display on separable modules means a cracked screen does not scrap a working processor, and "
        "favouring through-hole, widely stocked components keeps repair inside the community."
    ),
]

CONCLUSION = (
    "The upgrades requested here are blocked by SRAM capacity and by a single-threaded control flow, not by "
    "insufficient clock speed. Moving to a 32-bit RISC-V part with integrated Bluetooth Low Energy, placing "
    "deadline-critical code in deterministic on-chip memory while allowing cached flash execution elsewhere, "
    "and replacing the superloop with interrupt-driven acquisition into a batched ring buffer brings the "
    "daily cost to roughly 3.3 mAh against a 100 mAh budget. Headroom in the memory hierarchy is what makes "
    "incremental growth survivable, and is therefore the cheapest available defence against premature "
    "obsolescence."
)

# APA 7 reference entries as (author_and_date, italic_title, tail).
# The title element of a standalone work is italic sentence case (APA 7, Sec. 9.19).
# Where the group author is also the publisher, the publisher is omitted from the
# source element (APA 7, Sec. 9.29). Alphabetical, unnumbered, 0.5 in hanging indent.
REFERENCES = [
    (
        "Barry, R. (2016). ",
        "Mastering the FreeRTOS real time kernel: A hands-on tutorial guide",
        ". Real Time Engineers Ltd.",
    ),
    (
        "Bluetooth Special Interest Group. (2019). ",
        "Bluetooth core specification",
        " (Version 5.1).",
    ),
    (
        "Espressif Systems. (2023). ",
        "ESP32-C3 series datasheet",
        ". https://documentation.espressif.com/esp32-c3_datasheet_en.pdf",
    ),
    (
        "Forti, V., Bald\u00e9, C. P., Kuehr, R., & Bel, G. (2020). ",
        "The global e-waste monitor 2020: Quantities, flows and the circular economy potential",
        ". United Nations University; International Telecommunication Union; International Solid "
        "Waste Association.",
    ),
    (
        "Microchip Technology. (2016). ",
        "ATmega328/P datasheet",
        ".",
    ),
    (
        "Patterson, D. A., & Hennessy, J. L. (2021). ",
        "Computer organization and design: The hardware software interface, RISC-V edition",
        " (2nd ed.). Morgan Kaufmann.",
    ),
    (
        "SD Association. (2010). ",
        "SD specifications part 1: Physical layer simplified specification",
        " (Version 3.01).",
    ),
    (
        "Sensirion AG. (2019). ",
        "Datasheet SHT3x-DIS: Humidity and temperature sensor",
        ".",
    ),
    (
        "Sitronix Technology. (2011). ",
        "ST7735S 132 \u00d7 162 pixel TFT single chip driver datasheet",
        ".",
    ),
    (
        "Winbond Electronics. (2021). ",
        "W25Q64JV 3V 64M-bit serial flash memory datasheet",
        ".",
    ),
]


def set_base_style(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = st.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(3)

    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.name = "Times New Roman"


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11)
    r.font.name = "Times New Roman"


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(10)
    r.font.name = "Times New Roman"


def add_table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            if j == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(cell)
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)
            if i == 0 or row[0] == "Total":
                r.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def main():
    doc = Document()
    set_base_style(doc)

    add_title(doc, TITLE)

    add_heading(doc, "Introduction")
    add_body(doc, INTRO)

    add_heading(doc, "Background and Related Work")
    for para in BACKGROUND:
        add_body(doc, para)

    add_heading(doc, "Methodology and System Design")
    add_body(doc, METH_1)
    add_caption(doc, TABLE1_CAP)
    add_table(doc, TABLE1)
    add_body(doc, METH_2)
    add_body(doc, METH_3)
    add_body(doc, METH_4)
    add_body(doc, METH_5)
    add_body(doc, METH_6)
    add_caption(doc, TABLE2_CAP)
    add_table(doc, TABLE2)

    add_heading(doc, "Discussion")
    for para in DISCUSSION:
        add_body(doc, para)

    add_heading(doc, "Conclusion")
    add_body(doc, CONCLUSION)

    doc.add_page_break()
    add_heading(doc, "References")
    for head, title, tail in REFERENCES:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.space_after = Pt(3)
        for text, italic in ((head, False), (title, True), (tail, False)):
            r = p.add_run(text)
            r.italic = italic
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)

    target = OUT
    try:
        doc.save(target)
    except PermissionError:
        base, ext = os.path.splitext(OUT)
        n = 2
        while os.path.exists(f"{base}_v{n}{ext}"):
            n += 1
        target = f"{base}_v{n}{ext}"
        doc.save(target)
        print(f"NOTE: '{OUT}' is open in Word and locked, so it was not overwritten.")
        print(f"      Close it and delete the older copy once you have this one.")

    words = 0
    for chunk in [INTRO] + BACKGROUND + [METH_1, METH_2, METH_3, METH_4, METH_5, METH_6] \
                 + DISCUSSION + [CONCLUSION]:
        words += len(chunk.split())
    print(f"saved: {target}")
    print(f"body prose word count (excludes tables, headings, references): {words}")


if __name__ == "__main__":
    main()
